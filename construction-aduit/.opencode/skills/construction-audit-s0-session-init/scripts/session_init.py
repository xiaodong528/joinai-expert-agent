#!/usr/bin/env python3
"""
session_init.py - S0 session bootstrap for construction audit.

This script validates the uploaded rule document and spreadsheet, preserves the
real spreadsheet format in a working copy, and writes audit-config.yaml.
"""

from __future__ import annotations

import argparse
import shutil
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

try:
    import yaml
except ImportError:  # pragma: no cover - yaml is expected in normal environments
    yaml = None  # type: ignore

try:
    from docx import Document
except ImportError:  # pragma: no cover - python-docx is expected in normal environments
    Document = None  # type: ignore

try:
    import openpyxl
except ImportError:  # pragma: no cover - openpyxl is expected in normal environments
    openpyxl = None  # type: ignore

try:
    import xlrd
except ImportError:  # pragma: no cover - xlrd is expected in normal environments
    xlrd = None  # type: ignore


XLS_SIGNATURE = b"\xD0\xCF\x11\xE0\xA1\xB1\x1A\xE1"
ZIP_SIGNATURES = (b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08")


def friendly_error(message: str) -> int:
    print(f"Error: {message}", file=sys.stderr)
    return 1


def ensure_parent_dir(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)


def read_signature(path: Path, size: int = 8) -> bytes:
    with path.open("rb") as file_obj:
        return file_obj.read(size)


def detect_excel_format(path: Path) -> str:
    """Detect the actual Excel format from file bytes, not the filename suffix."""
    signature = read_signature(path)

    if signature.startswith(ZIP_SIGNATURES):
        if openpyxl is None:
            raise ValueError("缺少 openpyxl，无法读取 xlsx 文件")
        try:
            with path.open("rb") as file_obj:
                workbook = openpyxl.load_workbook(file_obj, read_only=True, data_only=True)
            workbook.close()
        except Exception as exc:  # pragma: no cover - exercised by CLI tests
            raise ValueError(f"无法读取 Excel 文件: {path}") from exc
        return "xlsx"

    if signature == XLS_SIGNATURE:
        if xlrd is None:
            raise ValueError("缺少 xlrd，无法读取 xls 文件")
        try:
            workbook = xlrd.open_workbook(file_contents=path.read_bytes())
            _ = workbook.sheet_names()
        except Exception as exc:  # pragma: no cover - exercised by CLI tests
            raise ValueError(f"无法读取 Excel 文件: {path}") from exc
        return "xls"

    raise ValueError(f"无法识别 Excel 文件格式: {path}")


def validate_rule_document(path: Path) -> Path:
    if not path.is_file():
        raise ValueError(f"规则文档不存在: {path}")
    if path.suffix.lower() != ".docx":
        raise ValueError(f"规则文档必须是 .docx 文件: {path}")
    if Document is None:
        raise ValueError("缺少 python-docx，无法读取规则文档")

    try:
        Document(str(path))
    except Exception as exc:  # pragma: no cover - exercised by CLI if malformed
        raise ValueError(f"无法读取规则文档: {path}") from exc
    return path.resolve()


def _classify_xlsx_sheet_visibility(path: Path) -> tuple[list[str], list[str], list[str]]:
    with path.open("rb") as file_obj:
        workbook = openpyxl.load_workbook(file_obj, read_only=False, data_only=True)
    try:
        all_sheets: list[str] = []
        visible_sheets: list[str] = []
        hidden_sheets: list[str] = []
        for worksheet in workbook.worksheets:
            all_sheets.append(worksheet.title)
            if getattr(worksheet, "sheet_state", "visible") == "visible":
                visible_sheets.append(worksheet.title)
            else:
                hidden_sheets.append(worksheet.title)
        return all_sheets, visible_sheets, hidden_sheets
    finally:
        workbook.close()


def _classify_xls_sheet_visibility(path: Path) -> tuple[list[str], list[str], list[str]]:
    workbook = xlrd.open_workbook(file_contents=path.read_bytes(), formatting_info=True)
    all_sheets = list(workbook.sheet_names())
    visibility = getattr(workbook, "_sheet_visibility", None)

    visible_sheets: list[str] = []
    hidden_sheets: list[str] = []
    for index, sheet_name in enumerate(all_sheets):
        is_hidden = (
            isinstance(visibility, (list, tuple))
            and index < len(visibility)
            and visibility[index] != 0
        )
        if is_hidden:
            hidden_sheets.append(sheet_name)
        else:
            visible_sheets.append(sheet_name)
    return all_sheets, visible_sheets, hidden_sheets


def load_sheet_metadata(path: Path, source_format: str) -> tuple[list[str], list[str], list[str]]:
    if source_format == "xlsx":
        return _classify_xlsx_sheet_visibility(path)
    return _classify_xls_sheet_visibility(path)


def normalize_workbook(
    original_path: Path,
    working_path: Path,
    source_format: str,
) -> Path:
    ensure_parent_dir(working_path)
    if working_path.suffix.lower() != f".{source_format}":
        raise ValueError(f"工作副本后缀必须与真实格式一致: {working_path}")
    if original_path.resolve() != working_path.resolve():
        shutil.copy2(original_path, working_path)
    return working_path.resolve()


def build_config(
    audit_type: str,
    project_name: str,
    rule_document_path: Path,
    spreadsheet_original_path: Path,
    spreadsheet_working_path: Path,
    sheet_scope: str,
    source_format: str,
    all_sheets: list[str],
    visible_sheets: list[str],
    hidden_sheets: list[str],
    output_dir: Path,
) -> dict[str, Any]:
    now = datetime.now()
    audit_id = now.strftime("AUDIT-%Y%m%d-%H%M%S")
    audit_date = now.date().isoformat()
    resolved_output_dir = output_dir.resolve()
    working_path_str = str(spreadsheet_working_path.resolve())
    target_sheets = list(visible_sheets)

    return {
        "audit_id": audit_id,
        "audit_type": audit_type,
        "audit_date": audit_date,
        "project_info": {
            "project_name": project_name,
        },
        "rule_document": {
            "path": str(rule_document_path.resolve()),
            "markdown_path": str((resolved_output_dir / "rule_doc.md").resolve()),
        },
        "spreadsheet": {
            "original_path": str(spreadsheet_original_path.resolve()),
            "working_path": working_path_str,
            "path": working_path_str,
            "source_format": source_format,
            "sheet_scope": sheet_scope,
            "sheets": target_sheets,
            "all_sheets": all_sheets,
            "visible_sheets": visible_sheets,
            "hidden_sheets": hidden_sheets,
        },
        "output_dir": str(resolved_output_dir),
    }


def write_config(config: dict[str, Any], output_path: Path) -> None:
    ensure_parent_dir(output_path)
    if yaml is not None:
        payload = yaml.safe_dump(config, allow_unicode=True, sort_keys=False)
    else:  # pragma: no cover - yaml should be installed
        import json

        payload = json.dumps(config, ensure_ascii=False, indent=2)
    output_path.write_text(payload, encoding="utf-8")


def emit_success_summary(config: dict[str, Any], config_path: Path) -> None:
    spreadsheet = config["spreadsheet"]
    summary = (
        f"audit_id={config['audit_id']} "
        f"source_format={spreadsheet['source_format']} "
        f"target_sheets_count={len(spreadsheet['sheets'])} "
        f"config_path={config_path.resolve()}"
    )
    print(summary)


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Bootstrap a construction audit session.")
    parser.add_argument("--rule-document", required=True, help="Path to rule_document.docx")
    parser.add_argument("--spreadsheet", required=True, help="Path to spreadsheet file")
    parser.add_argument(
        "--audit-type",
        required=True,
        choices=("budget", "settlement"),
        help="Audit type",
    )
    parser.add_argument(
        "--sheet-scope",
        default="visible",
        choices=("visible", "all"),
        help="Target sheet scope written into audit-config.yaml; formal flow only accepts 'visible'",
    )
    parser.add_argument("--project-name", required=True, help="Project name")
    parser.add_argument("--output-dir", required=True, help="Output directory")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)

    try:
        rule_document_path = validate_rule_document(Path(args.rule_document))
        spreadsheet_original_path = Path(args.spreadsheet)
        if not spreadsheet_original_path.is_file():
            return friendly_error(f"表格文件不存在: {spreadsheet_original_path}")

        source_format = detect_excel_format(spreadsheet_original_path)
        all_sheets, visible_sheets, hidden_sheets = load_sheet_metadata(
            spreadsheet_original_path,
            source_format,
        )
        if not all_sheets:
            return friendly_error("Excel 文件中未找到工作表")
        if args.sheet_scope != "visible":
            return friendly_error("当前正式主链只支持 sheet_scope=visible；hidden sheet 不能进入 spreadsheet.sheets")
        if not visible_sheets:
            return friendly_error("Excel 文件中未找到可直接审查的可见工作表")

        output_dir = Path(args.output_dir)
        output_dir.mkdir(parents=True, exist_ok=True)
        working_path = output_dir / f"{spreadsheet_original_path.stem}.{source_format}"
        normalize_workbook(spreadsheet_original_path, working_path, source_format)

        config = build_config(
            audit_type=args.audit_type,
            project_name=args.project_name,
            rule_document_path=rule_document_path,
            spreadsheet_original_path=spreadsheet_original_path,
            spreadsheet_working_path=working_path,
            sheet_scope=args.sheet_scope,
            source_format=source_format,
            all_sheets=all_sheets,
            visible_sheets=visible_sheets,
            hidden_sheets=hidden_sheets,
            output_dir=output_dir,
        )
        config_path = output_dir / "audit-config.yaml"
        write_config(config, config_path)
        emit_success_summary(config, config_path)
        return 0
    except ValueError as exc:
        return friendly_error(str(exc))


if __name__ == "__main__":
    raise SystemExit(main())
